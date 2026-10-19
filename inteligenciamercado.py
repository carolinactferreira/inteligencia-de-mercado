import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import pdfplumber
import re
import datetime

st.set_page_config(page_title="Relatório de Estoque", layout="wide")

st.title("📊 Gerador de Relatório de Estoque Imobiliário")
st.write("Faça o upload da planilha (.xlsx) ou arquivo PDF (.pdf) de estoque para gerar o relatório formatado.")

def converter_numero_br(val):
    if pd.isna(val) or val is None: return 0.0
    s = str(val).strip().replace('R$', '').replace('$', '').replace(' ', '')
    if not s or s.lower() in ['nan', 'none', 'null', '<na>', '-']: return 0.0
    if ',' in s and '.' in s: s = s.replace('.', '').replace(',', '.')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(re.sub(r'[^0-9.]', '', s))
    except: return 0.0

def limpar_texto(val, padrao='-'):
    if pd.isna(val) or val is None: return padrao
    s = str(val).strip()
    if not s or s.lower() in ['nan', 'none', 'null', '<na>', '(nenhuma / não existe)']: return padrao
    return s

def parse_concatenated_string(text_str):
    s = str(text_str).strip()
    apto_match = re.search(r'(?:APTO|APT|UNIDADE|UP|AP)\s*[:\.-]?\s*(\w+)', s, re.IGNORECASE)
    apto = apto_match.group(1) if apto_match else '-'
    tipo_match = re.search(r'\b(\d\s*Q\s*S|\d\s*QTS\s*S|\d\s*DORM\s*S|\d\s*Q|\d\s*QTS|\d\s*DORM|\d\s*D)\b', s, re.IGNORECASE)
    tipologia = tipo_match.group(1) if tipo_match else '-'
    clean_torre = re.split(r'\s*-\s*\d\s*Q|\s*-\s*\d\s*D|\s*-\s*APTO|\s*-\s*UNIDADE', s, flags=re.IGNORECASE)[0].strip()
    bloco = clean_torre if clean_torre else s
    return bloco, apto, tipologia

def mapear_tipologia(t):
    t_str = str(t).strip().upper()
    if not t_str or t_str in ['-', 'NAN', 'NONE']: return 'OUTRA TIPOLOGIA'
    has_suite = any(term in t_str for term in ['SUITE', 'SUÍTE', ' S', 'S ']) or t_str.endswith('S')
    if '3' in t_str: return '3 Q S' if has_suite else '3 Q'
    elif '2' in t_str: return '2 Q S' if has_suite else '2 Q'
    elif '1' in t_str: return '1 Q S' if has_suite else '1 Q'
    return 'OUTRA TIPOLOGIA'

def categorizar_vaga(row, col_moto='', col_carro=''):
    moto_val = str(row.get(col_moto, '')).strip().upper() if col_moto in row else ''
    carro_val = str(row.get(col_carro, '')).strip().upper() if col_carro in row else ''
    if moto_val in ['1', 'SIM', 'S', 'MOTO', 'TRUE', '1.0']: return 'Vaga Moto'
    elif carro_val in ['1', 'SIM', 'S', 'CARRO', 'TRUE', '1.0']: return 'Vaga Carro'
    if converter_numero_br(moto_val) > 0: return 'Vaga Moto'
    if converter_numero_br(carro_val) > 0: return 'Vaga Carro'
    return 'Sem Vaga'

def auto_detect_columns(df):
    mapping = {'empreendimento': None, 'bloco': None, 'apto': None, 'tipologia': None, 'area': None, 'valor_imovel': None, 'valor_avaliacao': None, 'carro': None, 'moto': None}
    for col in df.columns:
        c_lower = str(col).lower()
        if any(k in c_lower for k in ['empreend', 'projeto', 'condominio', 'condomínio', 'obra']) and not mapping['empreendimento']: mapping['empreendimento'] = col
        elif any(k in c_lower for k in ['bloco', 'torre']) and not mapping['bloco']: mapping['bloco'] = col
        elif any(k in c_lower for k in ['apto', 'unidade', 'apartamento']) and not mapping['apto']: mapping['apto'] = col
        elif any(k in c_lower for k in ['tipologia', 'tipo', 'planta', 'dorm']) and not mapping['tipologia']: mapping['tipologia'] = col
        elif any(k in c_lower for k in ['area', 'área', 'm2', 'privativa']) and not mapping['area']: mapping['area'] = col
        elif any(k in c_lower for k in ['venda', 'imovel', 'imóvel', 'preço', 'preco']) and 'aval' not in c_lower and not mapping['valor_imovel']: mapping['valor_imovel'] = col
        elif any(k in c_lower for k in ['aval']) and not mapping['valor_avaliacao']: mapping['valor_avaliacao'] = col
        elif 'carro' in c_lower and not mapping['carro']: mapping['carro'] = col
        elif 'moto' in c_lower and not mapping['moto']: mapping['moto'] = col

    sample_rows = df.head(10)
    for col in df.columns:
        vals = [str(v).strip() for v in sample_rows[col].dropna()]
        if not mapping['bloco'] and any('TORRE' in v.upper() or 'BLOCO' in v.upper() for v in vals) and any('APTO' in v.upper() for v in vals): mapping['bloco'] = col
        if not mapping['area']:
            numeric_vals = []
            for v in vals:
                try:
                    vf = float(v.replace(',', '.'))
                    if 10.0 <= vf <= 400.0: numeric_vals.append(vf)
                except: pass
            if len(numeric_vals) >= len(vals) * 0.4 and len(vals) > 0: mapping['area'] = col
    return mapping

def get_area_series(df, sel_area):
    if sel_area in df.columns and sel_area != "(Nenhuma / Não existe)": return df[sel_area]
    for col in df.columns:
        vals = df[col].dropna()
        parsed_vals = []
        for v in vals:
            try:
                vf = float(str(v).replace(',', '.'))
                if 15.0 <= vf <= 400.0: parsed_vals.append(vf)
            except: pass
        if len(parsed_vals) >= len(vals) * 0.4 and len(vals) > 0: return df[col]
    return pd.Series([0] * len(df))

def extrair_tabelas_pdf(file_bytes):
    todas_linhas = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            tabelas = page.extract_tables()
            for tabela in tabelas:
                if tabela:
                    for r in tabela:
                        if any(r): todas_linhas.append([str(c).strip().replace('\n', ' ') if c is not None else '' for c in r])
            if not tabelas:
                texto = page.extract_text()
                if texto:
                    for line in texto.split('\n'):
                        parts = line.split()
                        if len(parts) >= 3: todas_linhas.append(parts)
    if not todas_linhas: return pd.DataFrame()

    header_idx = 0
    for idx, row in enumerate(todas_linhas[:10]):
        row_str = " ".join(row).lower()
        if any(k in row_str for k in ['bloco', 'torre', 'apto', 'unidade', 'tipologia', 'valor', 'área', 'area', 'vaga']):
            header_idx = idx
            break
            
    header = todas_linhas[header_idx]
    seen = {}
    header_unico = []
    for h in header:
        clean_h = h if h else "col"
        if clean_h in seen:
            seen[clean_h] += 1
            header_unico.append(f"{clean_h}_{seen[clean_h]}")
        else:
            seen[clean_h] = 0
            header_unico.append(clean_h)
            
    dados = todas_linhas[header_idx + 1:]
    dados_limpos = []
    for r in dados:
        if len(r) == len(header_unico):
            if " ".join(r).lower() != " ".join(header).lower(): dados_limpos.append(r)
        elif len(r) < len(header_unico): dados_limpos.append(r + [''] * (len(header_unico) - len(r)))
        else: dados_limpos.append(r[:len(header_unico)])
            
    return pd.DataFrame(dados_limpos, columns=header_unico)

uploaded_file = st.file_uploader("Arraste sua planilha (.xlsx) ou arquivo (.pdf) aqui", type=["xlsx", "pdf"])

if uploaded_file is not None:
    file_bytes = uploaded_file.read()
    
    if uploaded_file.name.endswith('.pdf'): df = extrair_tabelas_pdf(file_bytes)
    else:
        df = pd.read_excel(io.BytesIO(file_bytes))
        df.columns = [str(c).strip().replace('\n', ' ') for c in df.columns]
        
    if df.empty:
        st.error("Não foi possível extrair dados do arquivo enviado.")
    else:
        st.subheader("📋 Dados Brutos Identificados no Arquivo")
        st.write(f"Total de registros encontrados: **{len(df)} linhas**")
        st.dataframe(df.head(15), use_container_width=True)

        st.subheader("⚙️ Mapeamento de Colunas")
        colunas_disponiveis = ["(Nenhuma / Não existe)"] + list(df.columns)
        mapping_auto = auto_detect_columns(df)

        def idx_of(val): return colunas_disponiveis.index(val) if val in colunas_disponiveis else 0

        col01, col02 = st.columns(2)
        with col01: sel_emp = st.selectbox("Coluna do Empreendimento:", colunas_disponiveis, index=idx_of(mapping_auto['empreendimento']))
        with col02: sel_bloco = st.selectbox("Coluna do Bloco/Torre:", colunas_disponiveis, index=idx_of(mapping_auto['bloco']))

        col1, col2, col3 = st.columns(3)
        with col1:
            sel_apto = st.selectbox("Coluna do Apto/Unidade:", colunas_disponiveis, index=idx_of(mapping_auto['apto']))
            sel_tipo = st.selectbox("Coluna da Tipologia/Planta:", colunas_disponiveis, index=idx_of(mapping_auto['tipologia']))
        with col2:
            sel_area = st.selectbox("Coluna da Área Total/Privativa:", colunas_disponiveis, index=idx_of(mapping_auto['area']))
            sel_valor = st.selectbox("Coluna do Valor do Imóvel:", colunas_disponiveis, index=idx_of(mapping_auto['valor_imovel']))
        with col3:
            sel_avaliacao = st.selectbox("Coluna do Valor da Avaliação:", colunas_disponiveis, index=idx_of(mapping_auto['valor_avaliacao']))
            sel_carro = st.selectbox("Coluna Vaga Carro:", colunas_disponiveis, index=idx_of(mapping_auto['carro']))

        sel_moto = st.selectbox("Coluna Vaga Moto:", colunas_disponiveis, index=idx_of(mapping_auto['moto']))

        st.markdown("---")

        lista_emps = []
        if sel_emp != "(Nenhuma / Não existe)":
            lista_emps = [str(x).strip() for x in df[sel_emp].dropna().unique() if str(x).strip() and str(x).lower() != 'nan']

        if lista_emps:
            st.subheader("🏢 Seleção de Empreendimento")
            filtro_emp = st.multiselect(
                "Selecione os empreendimentos para incluir no relatório:",
                options=lista_emps,
                default=lista_emps
            )
        else:
            filtro_emp = []

        if st.button("🚀 Gerar Relatório Formatado em Excel", type="primary"):
            with st.spinner("Processando dados e formatando Excel..."):
                
                if lista_emps and filtro_emp:
                    df_proc_source = df[df[sel_emp].astype(str).str.strip().isin(filtro_emp)].copy()
                    nome_relatorio = filtro_emp[0] if len(filtro_emp) == 1 else "Varios_Empreendimentos"
                else:
                    df_proc_source = df.copy()
                    nome_relatorio = "Relatorio_Geral"

                s_emp = df_proc_source[sel_emp] if sel_emp in df_proc_source.columns else pd.Series(['-']*len(df_proc_source))
                s_bloco = df_proc_source[sel_bloco] if sel_bloco in df_proc_source.columns else pd.Series(['-']*len(df_proc_source))
                s_apto = df_proc_source[sel_apto] if sel_apto in df_proc_source.columns else pd.Series(['-']*len(df_proc_source))
                s_tipo = df_proc_source[sel_tipo] if sel_tipo in df_proc_source.columns else pd.Series(['-']*len(df_proc_source))
                s_area = get_area_series(df_proc_source, sel_area)
                s_valor = df_proc_source[sel_valor] if sel_valor in df_proc_source.columns else pd.Series([0]*len(df_proc_source))
                s_aval = df_proc_source[sel_avaliacao] if sel_avaliacao in df_proc_source.columns else pd.Series([0]*len(df_proc_source))

                res_emp, res_bloco, res_apto, res_tipo_orig = [], [], [], []

                for idx in range(len(df_proc_source)):
                    raw_e = str(s_emp.iloc[idx]).strip()
                    raw_b = str(s_bloco.iloc[idx]).strip()
                    raw_a = str(s_apto.iloc[idx]).strip()
                    raw_t = str(s_tipo.iloc[idx]).strip()

                    if ('TORRE' in raw_b.upper() or 'BLOCO' in raw_b.upper()) and ('APTO' in raw_b.upper() or 'Q' in raw_b.upper()):
                        ext_b, ext_a, ext_t = parse_concatenated_string(raw_b)
                        final_b = ext_b if ext_b else raw_b
                        final_a = raw_a if raw_a not in ['-', 'nan', 'None', '', '(Nenhuma / Não existe)'] else ext_a
                        final_t = raw_t if raw_t not in ['-', 'nan', 'None', '', '(Nenhuma / Não existe)'] else ext_t
                    else:
                        final_b = raw_b if raw_b not in ['nan', 'None', ''] else '-'
                        final_a = raw_a if raw_a not in ['nan', 'None', ''] else '-'
                        final_t = raw_t if raw_t not in ['nan', 'None', ''] else '-'

                    res_emp.append(limpar_texto(raw_e))
                    res_bloco.append(limpar_texto(final_b))
                    res_apto.append(limpar_texto(final_a))
                    res_tipo_orig.append(limpar_texto(final_t))

                df_proc = pd.DataFrame({
                    'Empreendimento': res_emp,
                    'Bloco': res_bloco,
                    'Apto': res_apto,
                    'Tipologia_Orig': res_tipo_orig,
                    'Área': [converter_numero_br(v) for v in s_area],
                    'Valor_Imovel': [converter_numero_br(v) for v in s_valor],
                    'Valor_Avaliacao': [converter_numero_br(v) for v in s_aval],
                    'Categoria_Vaga': [categorizar_vaga(row, col_moto=sel_moto, col_carro=sel_carro) for _, row in df_proc_source.iterrows()]
                })
                df_proc['Tipologia_Clean'] = df_proc['Tipologia_Orig'].apply(mapear_tipologia)

                output = io.BytesIO()
                wb = openpyxl.Workbook()
                wb.remove(wb.active)

                data_hoje = datetime.datetime.now().strftime("%d/%m")
                fill_title = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
                font_title = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
                font_subtitle = Font(name='Calibri', size=11, italic=True, color='FFFFFF')
                fill_kpi_header = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid')
                font_kpi_header = Font(name='Calibri', size=10, bold=True, color='595959')
                fill_kpi_val = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
                font_kpi_val = Font(name='Calibri', size=16, bold=True, color='1B365D')
                font_kpi_sub = Font(name='Calibri', size=9, italic=True, color='808080')
                border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

                headers = ['Empreendimento', 'Bloco/Torre', 'Apto/Unidade', 'Tipologia Original', 'Área Total (m²)', 'Valor Imóvel (R$)', 'Valor Avaliação (R$)']

                def aplicar_kpis(ws, titulo, df_tab):
                    total_unid = len(df_tab)
                    area_media = df_tab['Área'].mean() if total_unid > 0 else 0
                    val_imovel = df_tab['Valor_Imovel'].mean() if total_unid > 0 else 0
                    val_aval = df_tab['Valor_Avaliacao'].mean() if total_unid > 0 else 0

                    ws.merge_cells('A1:G1')
                    ws['A1'] = titulo
                    ws['A1'].font = font_title
                    ws['A1'].fill = fill_title
                    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

                    ws.merge_cells('A2:G2')
                    ws['A2'] = f"Resumo Automático de Estoque e Vagas (Base: {data_hoje})"
                    ws['A2'].font = font_subtitle
                    ws['A2'].fill = fill_title
                    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

                    ws.merge_cells('A4:B4')
                    ws['A4'] = "TOTAL DE UNIDADES"
                    ws.merge_cells('C4:D4')
                    ws['C4'] = "ÁREA MÉDIA (TOTAL)"
                    ws['E4'] = "VALOR MÉDIO DO IMÓVEL"
                    ws.merge_cells('F4:G4')
                    ws['F4'] = "VALOR MÉDIO DE AVALIAÇÃO"

                    for c in ['A4','B4','C4','D4','E4','F4','G4']:
                        ws[c].fill = fill_kpi_header
                        ws[c].font = font_kpi_header
                        ws[c].alignment = Alignment(horizontal='center', vertical='center')
                        ws[c].border = border_thin

                    ws.merge_cells('A5:B5')
                    ws['A5'] = total_unid
                    ws.merge_cells('C5:D5')
                    ws['C5'] = f"{area_media:,.2f} m²".replace(',', 'X').replace('.', ',').replace('X', '.')
                    ws['E5'] = f"R$ {val_imovel:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
                    ws.merge_cells('F5:G5')
                    ws['F5'] = f"R$ {val_aval:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

                    for c in ['A5','B5','C5','D5','E5','F5','G5']:
                        ws[c].fill = fill_kpi_val
                        ws[c].font = font_kpi_val
                        ws[c].alignment = Alignment(horizontal='center', vertical='center')
                        ws[c].border = border_thin

                    ws.merge_cells('A6:B6')
                    ws['A6'] = "unidades no estoque"
                    ws.merge_cells('C6:D6')
                    ws['C6'] = "Metragem média calculada"
                    ws['E6'] = "Valor médio de tabela"
                    ws.merge_cells('F6:G6')
                    ws['F6'] = "Valor médio de avaliação"

                    for c in ['A6','B6','C6','D6','E6','F6','G6']:
                        ws[c].fill = fill_kpi_val # Usa a mesma cor de fundo
                        ws[c].font = font_kpi_sub
                        ws[c].alignment = Alignment(horizontal='center', vertical='center')
                        ws[c].border = border_thin

                    ws.append([])
                    ws.append(headers)

                    for _, row in df_tab.iterrows():
                        ws.append([row['Empreendimento'], row['Bloco'], row['Apto'], row['Tipologia_Orig'], row['Área'], row['Valor_Imovel'], row['Valor_Avaliacao']])

                ws_geral = wb.create_sheet(title="Estoque Geral")
                aplicar_kpis(ws_geral, f"RELATÓRIO DE ESTOQUE — {nome_relatorio.upper()}", df_proc)

                configuracoes = [
                    ('Moto 1Q', 'Vaga Moto', '1 Q'), ('Carro 1Q', 'Vaga Carro', '1 Q'), ('Sem Vaga 1Q', 'Sem Vaga', '1 Q'),
                    ('Moto 2QS', 'Vaga Moto', '2 Q S'), ('Carro 2QS', 'Vaga Carro', '2 Q S'), ('Sem Vaga 2QS', 'Sem Vaga', '2 Q S'),
                    ('Moto 2Q', 'Vaga Moto', '2 Q'), ('Carro 2Q', 'Vaga Carro', '2 Q'), ('Sem Vaga 2Q', 'Sem Vaga', '2 Q'),
                    ('Outras Unidades', 'TODAS', 'OUTRA TIPOLOGIA')
                ]

                for sheet_name, cat_key, tipo_key in configuracoes:
                    if cat_key == 'TODAS': sub_df = df_proc[df_proc['Tipologia_Clean'] == tipo_key]
                    else: sub_df = df_proc[(df_proc['Categoria_Vaga'] == cat_key) & (df_proc['Tipologia_Clean'] == tipo_key)]
                    
                    if sub_df.empty: continue
                    
                    ws = wb.create_sheet(title=sheet_name)
                    aplicar_kpis(ws, f"{nome_relatorio.upper()} — {cat_key.upper()} ({tipo_key})", sub_df)

                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(min_row=9, max_col=7): # Dados começam agora na linha 9!
                        if len(row) >= 7:
                            row[4].number_format = '#,##0.00'
                            row[5].number_format = 'R$ #,##0.00'
                            row[6].number_format = 'R$ #,##0.00'
                    for col in sheet.columns:
                        max_len = max(len(str(cell.value or '')) for cell in col)
                        col_letter = get_column_letter(col[0].column)
                        sheet.column_dimensions[col_letter].width = max(max_len + 3, 15)

                wb.save(output)
                processed_data = output.getvalue()
                
                st.success("Relatório gerado com sucesso!")
                
                st.download_button(
                    label="📥 Baixar Relatório Excel Formatado",
                    data=processed_data,
                    file_name=f"Relatorio_{nome_relatorio.replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )  